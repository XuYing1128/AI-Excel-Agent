"""Native Excel chart helpers built on openpyxl."""

from __future__ import annotations

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.worksheet.worksheet import Worksheet


def add_bar_chart(
    ws: Worksheet,
    title: str,
    data_min_col: int,
    data_max_col: int,
    data_min_row: int,
    data_max_row: int,
    cats_col: int,
    cats_min_row: int,
    cats_max_row: int,
    anchor: str = "D3",
) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = "金额/数量"
    chart.x_axis.title = "分类"
    data = Reference(ws, min_col=data_min_col, max_col=data_max_col, min_row=data_min_row, max_row=data_max_row)
    cats = Reference(ws, min_col=cats_col, min_row=cats_min_row, max_row=cats_max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 12
    # These charts reference summary cells that are Excel formulas. Force a full
    # recalculation when the file is opened so Excel/WPS compute those values and
    # the chart is not blank.
    ws.parent.calculation.fullCalcOnLoad = True
    ws.add_chart(chart, anchor)


def add_line_chart(
    ws: Worksheet,
    title: str,
    data_min_col: int,
    data_max_col: int,
    data_min_row: int,
    data_max_row: int,
    cats_col: int,
    cats_min_row: int,
    cats_max_row: int,
    anchor: str = "D15",
) -> None:
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = "趋势"
    chart.x_axis.title = "期间"
    data = Reference(ws, min_col=data_min_col, max_col=data_max_col, min_row=data_min_row, max_row=data_max_row)
    cats = Reference(ws, min_col=cats_col, min_row=cats_min_row, max_row=cats_max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 12
    # These charts reference summary cells that are Excel formulas. Force a full
    # recalculation when the file is opened so Excel/WPS compute those values and
    # the chart is not blank.
    ws.parent.calculation.fullCalcOnLoad = True
    ws.add_chart(chart, anchor)


def add_pie_chart(
    ws: Worksheet,
    title: str,
    data_col: int,
    data_min_row: int,
    data_max_row: int,
    cats_col: int,
    cats_min_row: int,
    cats_max_row: int,
    anchor: str = "D3",
) -> None:
    chart = PieChart()
    chart.title = title
    data = Reference(ws, min_col=data_col, min_row=data_min_row, max_row=data_max_row)
    cats = Reference(ws, min_col=cats_col, min_row=cats_min_row, max_row=cats_max_row)
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 9
    # These charts reference summary cells that are Excel formulas. Force a full
    # recalculation when the file is opened so Excel/WPS compute those values and
    # the chart is not blank.
    ws.parent.calculation.fullCalcOnLoad = True
    ws.add_chart(chart, anchor)

